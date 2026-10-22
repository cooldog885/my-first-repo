#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将抓取得到的 Markdown 文件合并导出为单个 Word 文档（.docx）

支持轻量 Markdown 语法转换：
    # ~ ######  标题
    >            引用
    - / *        无序列表
    ---          分隔线
    其余         普通段落
"""

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_horizontal_rule(paragraph):
    """在段落下方添加一条水平分隔线"""
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def _parse_and_append(doc, md_text, page_break_before=False):
    """
    将一段 Markdown 文本解析并追加到 doc 中

    Args:
        doc: Document 对象
        md_text: Markdown 文本
        page_break_before: 是否在内容前插入分页符
    """
    if page_break_before:
        doc.add_page_break()

    blocks = []
    current = []
    for line in md_text.splitlines():
        if line.strip() == '':
            if current:
                blocks.append(current)
                current = []
        else:
            current.append(line)
    if current:
        blocks.append(current)

    for block in blocks:
        # 标题
        first = block[0].lstrip()
        if first.startswith('#'):
            level = 0
            for ch in first:
                if ch == '#':
                    level += 1
                else:
                    break
            level = min(max(level, 1), 6)
            text = first[level:].strip()
            doc.add_heading(text, level=level)
            continue

        # 分隔线
        if all(seg.strip() in ('---', '***', '___') for seg in block):
            p = doc.add_paragraph()
            _add_horizontal_rule(p)
            continue

        # 引用块
        if all(seg.lstrip().startswith('>') for seg in block):
            quote_text = '\n'.join(seg.lstrip()[1:].strip() for seg in block)
            try:
                p = doc.add_paragraph(quote_text, style='Quote')
            except KeyError:
                p = doc.add_paragraph(quote_text)
                for run in p.runs:
                    run.italic = True
            continue

        # 无序列表
        if all(seg.lstrip().startswith(('-', '*', '+')) for seg in block):
            for seg in block:
                item = seg.lstrip()[1:].strip()
                doc.add_paragraph(item, style='List Bullet')
            continue

        # 普通段落（多行合并）
        paragraph_text = '\n'.join(block)
        doc.add_paragraph(paragraph_text)


def merge_markdown_to_docx(md_files, output_path, title=None,
                           page_break_between=True):
    """
    将多个 Markdown 文件合并为一个 Word 文档

    Args:
        md_files: Markdown 文件路径列表（按期望顺序）
        output_path: 输出的 .docx 路径
        title: 文档总标题（可选）
        page_break_between: 每个文件之间是否分页

    Returns:
        int: 成功合并的文件数量
    """
    doc = Document()

    if title:
        doc.add_heading(title, level=0)

    count = 0
    for idx, md_path in enumerate(md_files):
        if not os.path.isfile(md_path):
            continue
        try:
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except Exception:
            continue

        if not content.strip():
            continue

        _parse_and_append(
            doc, content,
            page_break_before=(page_break_between and idx > 0)
        )
        count += 1

    doc.save(output_path)
    return count


def collect_markdown_files(root_path):
    """
    递归收集目录下的所有 .md 文件，按路径排序

    Args:
        root_path: 根目录

    Returns:
        list: 排序后的 .md 文件路径列表
    """
    result = []
    for dirpath, _, filenames in os.walk(root_path):
        for name in filenames:
            if name.lower().endswith('.md'):
                result.append(os.path.join(dirpath, name))
    result.sort()
    return result
